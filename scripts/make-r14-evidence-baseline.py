"""Prepare compact field/row provenance once; no original table is installed."""
from pathlib import Path
import csv, hashlib, sqlite3, sys
from openpyxl import load_workbook

source, target = map(Path, sys.argv[1:3])
target.parent.mkdir(parents=True, exist_ok=True)
if target.exists(): target.unlink()
db=sqlite3.connect(target)
db.executescript("CREATE TABLE facility_evidence(name TEXT,region TEXT,room TEXT,source TEXT,rownum INTEGER,field TEXT,sheet TEXT,PRIMARY KEY(name,source,rownum)) WITHOUT ROWID;CREATE TABLE evidence_sources(source TEXT PRIMARY KEY,signature TEXT);")
count=0
for path in sorted(source.rglob('*')):
    if path.suffix.lower() not in ['.csv','.xlsx'] or path.name.startswith('中继段') or path.name=='集群管理.xlsx':continue
    worksheet=''
    if path.suffix.lower()=='.csv':
        stream=path.open(encoding='gb18030',newline='');rows=csv.reader(stream)
    else:
        book=load_workbook(path,read_only=True,data_only=True);sheet=book.worksheets[0];worksheet=sheet.title;rows=sheet.iter_rows(values_only=True)
    headers=[str(v or '').strip() for v in next(rows)]
    def index(names):return next((i for i,h in enumerate(headers) if h in names),-1)
    name=index(['设施名称','节点名称','设备名称']);region=index(['所属区县','区域','所属区域']);room=index(['机房名称','所属机房'])
    if name<0:continue
    data=[]
    for row_number,row in enumerate(rows,2):
        def value(i):return str(row[i] or '').strip() if 0<=i<len(row) else ''
        if not value(name):continue
        field=chr(65+room)+'列 · '+headers[room] if 0<=room<26 else headers[room] if room>=0 else ''
        data.append((value(name),value(region),value(room),path.name,row_number,field,worksheet))
        if len(data)>=2000:db.executemany('INSERT OR REPLACE INTO facility_evidence VALUES(?,?,?,?,?,?,?)',data);count+=len(data);data=[]
    if data:db.executemany('INSERT OR REPLACE INTO facility_evidence VALUES(?,?,?,?,?,?,?)',data);count+=len(data)
    signature=hashlib.sha256(path.read_bytes()).hexdigest().upper()+'|'+worksheet
    db.execute('INSERT INTO evidence_sources VALUES(?,?)',(path.name,signature));db.commit()
    if path.suffix.lower()=='.csv':stream.close()
    else:book.close()
db.execute('VACUUM');db.close()
print('EVIDENCE BASELINE',count,'rows',target.stat().st_size,'bytes',target)
